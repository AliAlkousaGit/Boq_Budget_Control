# Copyright (c) 2026, BOQ Budget Control and contributors
# For license information, please see license.txt
"""BOQ Excel import — pure parser (no Frappe imports, unit-testable).

``build_preview(path, masters, uoms, sheet=None)`` reads a client / estimator BOQ
workbook and returns a preview dict: detected sheet + header, column mapping,
categories (matched to BOQ Budget Category masters, >= 80% similarity) with their
lines, skipped rows, sheet totals, warnings and blocking errors.

The Frappe side (``api.preview_boq_import``) supplies the masters / UOMs and the
form script turns an error-free preview into Category + Item rows.
"""

import re
from difflib import SequenceMatcher

import openpyxl

MATCH_THRESHOLD = 0.80
HEADER_SCAN_ROWS = 50
EMPTY_ROWS_TO_STOP = 300
UNPRICED_REMARK = "Unpriced in source BOQ"

# header synonyms, compared on _key() (lower-case alphanumerics only)
HEADER_SYNONYMS = {
	"number": {"sno", "slno", "srno", "serialno", "itemno", "no", "sn", "sl", "ref", "refno", "item"},
	"description": {"description", "desc", "itemdescription", "particulars", "workdescription",
					"descriptionofwork", "descriptionofworks", "scopeofwork"},
	"unit": {"unit", "units", "uom"},
	"qty": {"qty", "quantity", "qnty", "qtys", "quantities"},
	"rate": {"rate", "unitrate", "unitprice", "price", "rateaed", "unitrateaed"},
	"amount": {"amount", "amountaed", "amt", "totalamount", "value"},
	"total": {"total", "totalaed", "subtotal"},
	"remarks": {"remarks", "remark", "comments", "comment", "notes"},
}

UOM_ALIASES = {
	"ls": ["ls", "lumpsum", "lot"],
	"nos": ["nos", "no", "number", "numbers", "unit", "units", "each", "ea", "pcs", "pc"],
	"m2": ["m2", "sqm", "sqmt", "squaremeter", "squaremetre"],
	"m3": ["m3", "cum", "cbm", "cubicmeter", "cubicmetre"],
	"lm": ["lm", "rm", "rmt", "lmt", "m", "meter", "metre", "runningmeter", "linearmeter"],
	"kg": ["kg", "kgs", "kilogram"],
	"ton": ["ton", "tons", "tonne", "mt"],
	"set": ["set", "sets"],
}

RE_SUBTOTAL = re.compile(r"sub\s*-?\s*total", re.I)
RE_TOTAL = re.compile(r"\b(grand\s*total|total|discount|vat)\b", re.I)
RE_STOP = re.compile(r"^\s*(terms\s*(and|&)\s*conditions|notes?\s*[:\n]|notes?\s*$)", re.I)
RE_SKIP_HEADING = re.compile(r"^\s*(inclusions?|exclusions?)\s*[:\-]?\s*$", re.I)


# ---------------------------------------------------------------------------
# small helpers
# ---------------------------------------------------------------------------
def _key(value):
	return re.sub(r"[^a-z0-9]", "", str(value or "").lower())


def _clean(value):
	"""Whitespace-only strings -> None; strings trimmed; numbers untouched."""
	if isinstance(value, str):
		value = value.strip()
		return value or None
	return value


def _num(value):
	"""Number from a cell (int/float or numeric text like '1,200.50'); else None."""
	value = _clean(value)
	if value is None or isinstance(value, bool):
		return None
	if isinstance(value, (int, float)):
		return float(value)
	text = str(value).replace(",", "").strip()
	try:
		return float(text)
	except ValueError:
		return None


def _is_text(value):
	value = _clean(value)
	return value is not None and _num(value) is None


def format_ref(value):
	"""S.No as text: 2.1 -> '2.10', 33.1200000000001 -> '33.12', 8.014 -> '8.014', 3 -> '3'."""
	n = _num(value)
	if n is None:
		return (_clean(value) or "") if value is not None else ""
	n = round(n, 6)
	if float(n).is_integer():
		return str(int(n))
	if round(n, 2) == n:
		return f"{n:.2f}"
	return f"{n:.6f}".rstrip("0")


def _is_integer_ref(value):
	n = _num(value)
	return n is not None and n >= 1 and float(round(n, 6)).is_integer()


def normalize_name(text):
	text = str(text or "").lower().replace("&", " and ")
	text = re.sub(r"[^a-z0-9]+", " ", text)
	return re.sub(r"\s+", " ", text).strip()


def _word_overlap(a_words, b_words):
	"""Dice coefficient on words; two words count as equal when >= 85% alike (typos)."""
	if not a_words or not b_words:
		return 0.0
	remaining = list(b_words)
	hits = 0
	for w in a_words:
		for i, o in enumerate(remaining):
			if w == o or SequenceMatcher(None, w, o).ratio() >= 0.85:
				hits += 1
				remaining.pop(i)
				break
	return 2.0 * hits / (len(a_words) + len(b_words))


def similarity(a, b):
	"""0..1 score: half character similarity, half word overlap.

	Character similarity alone rates 'External Works' vs 'Metal works' at 80%;
	requiring the words to agree as well removes such false matches.
	"""
	a, b = normalize_name(a), normalize_name(b)
	if not a or not b:
		return 0.0
	if a == b or a.replace(" ", "") == b.replace(" ", ""):
		return 1.0  # also 'Blockwork' == 'Block Work', 'Waterproofing' == 'Water Proofing'
	chars = max(
		SequenceMatcher(None, a, b).ratio(),
		SequenceMatcher(None, " ".join(sorted(a.split())), " ".join(sorted(b.split()))).ratio(),
	)
	return (chars + _word_overlap(a.split(), b.split())) / 2.0


def match_category(source_name, masters):
	"""Return (best, candidates) where best is a dict or None (< threshold)."""
	scored = []
	for m in masters:
		score = max(similarity(source_name, m.get("category_name")), similarity(source_name, m.get("name")))
		scored.append({"name": m["name"], "category_name": m.get("category_name") or m["name"],
					   "score": round(score, 3)})
	scored.sort(key=lambda x: x["score"], reverse=True)
	best = scored[0] if scored and scored[0]["score"] >= MATCH_THRESHOLD else None
	return best, scored[:3]


def match_uom(unit, uoms):
	"""Existing UOM name for a source unit text, or None."""
	k = _key(unit)
	if not k:
		return None
	by_key = {}
	for u in uoms:
		by_key.setdefault(_key(u), u)
	if k in by_key:
		return by_key[k]
	for aliases in UOM_ALIASES.values():
		if k in aliases:
			for alias in aliases:
				if alias in by_key:
					return by_key[alias]
	return None


# ---------------------------------------------------------------------------
# sheet / header detection
# ---------------------------------------------------------------------------
def _header_role(cell_value):
	k = _key(cell_value)
	if not k:
		return None
	for role, names in HEADER_SYNONYMS.items():
		if k in names:
			return role
	if "description" in k:
		return "description"
	return None


def find_header(rows):
	"""rows: list of tuples (first HEADER_SCAN_ROWS rows). Returns (row_index0, {role: col0}) or None."""
	for i, row in enumerate(rows[:HEADER_SCAN_ROWS]):
		mapping = {}
		for c, value in enumerate(row):
			role = _header_role(value)
			if role and role not in mapping:
				mapping[role] = c
		priced = sum(1 for r in ("unit", "qty", "rate", "amount") if r in mapping)
		if "description" in mapping and priced >= 2:
			return i, mapping
	return None


def _col_letter(c0):
	return openpyxl.utils.get_column_letter(c0 + 1)


def _sheet_rows(ws):
	"""All rows as value tuples, stopping after a long run of empty rows."""
	rows, empty_run = [], 0
	for row in ws.iter_rows(values_only=True):
		if any(_clean(v) is not None for v in row):
			empty_run = 0
		else:
			empty_run += 1
			if empty_run >= EMPTY_ROWS_TO_STOP:
				break
		rows.append(row)
	while rows and not any(_clean(v) is not None for v in rows[-1]):
		rows.pop()
	return rows


# ---------------------------------------------------------------------------
# main entry
# ---------------------------------------------------------------------------
def build_preview(path, masters, uoms, sheet=None):
	wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
	try:
		candidates = []
		for ws in wb.worksheets:
			if getattr(ws, "sheet_state", "visible") != "visible":
				continue
			head = []
			for row in ws.iter_rows(max_row=HEADER_SCAN_ROWS, values_only=True):
				head.append(row)
			if find_header(head):
				candidates.append(ws.title)

		result = {"sheets": candidates, "sheet": None, "errors": [], "warnings": [],
				  "categories": [], "skipped": [], "totals": []}
		if not candidates:
			result["errors"].append("No BOQ header row found (needs Description plus at least two of "
									"Unit / Qty / Rate / Amount within the first 50 rows).")
			return result
		if sheet and sheet not in candidates:
			result["errors"].append(f"Sheet '{sheet}' has no BOQ header row.")
			return result
		if not sheet:
			if len(candidates) > 1:
				result["needs_sheet"] = True
				return result
			sheet = candidates[0]

		rows = _sheet_rows(wb[sheet])
	finally:
		wb.close()

	result["sheet"] = sheet
	parse_rows(rows, masters, uoms, result)
	return result


def parse_rows(rows, masters, uoms, result):
	header = find_header(rows)
	h_idx, cols = header
	result["header_row"] = h_idx + 1
	result["columns"] = {role: _col_letter(c) for role, c in cols.items()}

	def cell(row, role):
		c = cols.get(role)
		return row[c] if c is not None and c < len(row) else None

	categories, skipped, totals, errors, warnings = [], [], [], [], []
	current = None
	skip_block = False

	for r0 in range(h_idx + 1, len(rows)):
		row, rownum = rows[r0], r0 + 1
		desc = _clean(cell(row, "description"))
		ref_raw = _clean(cell(row, "number"))
		if desc is None and _is_text(ref_raw):
			# label written in the S.No column (e.g. merged "GRAND TOTAL AMOUNT" cell)
			desc, ref_raw = ref_raw, None
		qty_raw, rate_raw = _clean(cell(row, "qty")), _clean(cell(row, "rate"))
		qty, rate = _num(qty_raw), _num(rate_raw)
		amount, total = _num(cell(row, "amount")), _num(cell(row, "total"))
		unit = _clean(cell(row, "unit"))
		remarks = _clean(cell(row, "remarks"))
		desc_text = str(desc) if desc is not None else ""

		if desc is None:
			# value-only row under a block = that block's sheet subtotal
			value = total if total is not None else amount
			if value is not None and ref_raw is None and qty is None and rate is None:
				if current is not None and current["sheet_subtotal"] is None:
					current["sheet_subtotal"] = value
				skipped.append({"row": rownum, "reason": "Subtotal (value only)", "text": str(value)})
			continue

		if categories and RE_STOP.match(desc_text):
			skipped.append({"row": rownum, "reason": "End of BOQ (notes / terms)", "text": desc_text[:80]})
			break

		if RE_SUBTOTAL.search(desc_text) and qty is None and rate is None:
			value = total if total is not None else amount
			if current is not None and value is not None and current["sheet_subtotal"] is None:
				current["sheet_subtotal"] = value
			skipped.append({"row": rownum, "reason": "Subtotal", "text": desc_text[:80]})
			continue

		if RE_TOTAL.search(desc_text) and qty is None and rate is None and ref_raw is None:
			value = total if total is not None else amount
			totals.append({"row": rownum, "label": desc_text[:80], "value": value})
			skipped.append({"row": rownum, "reason": "Total / discount", "text": desc_text[:80]})
			continue

		if RE_SKIP_HEADING.match(desc_text):
			skip_block = True
			skipped.append({"row": rownum, "reason": "Inclusion / exclusion heading", "text": desc_text[:80]})
			continue

		if _is_text(qty_raw):
			skipped.append({"row": rownum, "reason": f"Qty is text ('{str(qty_raw)[:40]}')", "text": desc_text[:80]})
			continue

		has_numbers = any(v is not None for v in (qty, rate, amount))

		if skip_block and ref_raw is None:
			skipped.append({"row": rownum, "reason": "Inclusion / exclusion text", "text": desc_text[:80]})
			continue
		skip_block = False

		if _is_integer_ref(ref_raw):
			current = _new_category(desc_text, ref_raw, rownum)
			categories.append(current)
			if qty is not None or rate is not None:
				# category row that carries its own qty / rate -> category + one line
				current["lines"].append(_new_line(desc_text, ref_raw, rownum, unit, qty, rate,
												  amount, remarks, uoms, warnings))
			continue

		if ref_raw is None and not has_numbers:
			skipped.append({"row": rownum, "reason": "Text row (no S.No, no qty/rate)", "text": desc_text[:80]})
			continue

		if current is None:
			errors.append(f"Row {rownum}: line '{desc_text[:60]}' appears before the first category.")
			continue
		current["lines"].append(_new_line(desc_text, ref_raw, rownum, unit, qty, rate,
										  amount, remarks, uoms, warnings))

	if not categories:
		errors.append("No category found (a category row needs a whole-number S.No and a description).")

	_match_and_check(categories, masters, errors, warnings)

	for cat in categories:
		cat["amount"] = round(sum(l["amount"] for l in cat["lines"]), 2)
		if cat["sheet_subtotal"] is not None and abs(cat["sheet_subtotal"] - cat["amount"]) > 0.01:
			warnings.append(
				f"'{cat['source_name']}' (row {cat['row']}): sheet subtotal {cat['sheet_subtotal']:,.2f} "
				f"≠ lines {cat['amount']:,.2f} (diff {cat['amount'] - cat['sheet_subtotal']:,.2f})."
			)
		if not cat["amount"]:
			warnings.append(f"'{cat['source_name']}' has 0 budget — POs under Stop will be blocked until priced.")

	grand = round(sum(c["amount"] for c in categories), 2)
	result.update({
		"categories": categories, "skipped": skipped, "totals": totals,
		"errors": result["errors"] + errors, "warnings": result["warnings"] + warnings,
		"summary": {"categories": len(categories),
					"lines": sum(len(c["lines"]) for c in categories),
					"unpriced_lines": sum(1 for c in categories for l in c["lines"] if l["unpriced"]),
					"total": grand},
	})
	return result


def _new_category(name, ref, rownum):
	return {"source_name": name, "boq_ref": format_ref(ref), "row": rownum, "lines": [],
			"sheet_subtotal": None, "match": None, "candidates": []}


def _new_line(desc, ref, rownum, unit, qty, rate, amount, remarks, uoms, warnings):
	notes = []
	unpriced = not qty or not rate
	if unpriced:
		notes.append(UNPRICED_REMARK)
	uom = match_uom(unit, uoms)
	if unit and not uom:
		notes.append(f"Unit: {unit}")
	if remarks:
		notes.append(str(remarks))
	q = qty if qty and qty > 0 else 1.0
	r = 0.0 if unpriced else rate
	line_amount = round(q * r, 2)
	if not unpriced and amount is not None and abs(amount - line_amount) > 0.01:
		warnings.append(f"Row {rownum}: sheet amount {amount:,.2f} ≠ qty × rate {line_amount:,.2f}; "
						"budget uses qty × rate.")
	return {"row": rownum, "boq_ref": format_ref(ref) if ref is not None else "", "description": desc,
			"unit": unit, "uom": uom, "qty": q, "rate": r, "amount": line_amount,
			"sheet_amount": amount, "remarks": " | ".join(notes) or None, "unpriced": unpriced}


def _match_and_check(categories, masters, errors, warnings):
	resolved = {}
	for cat in categories:
		best, candidates = match_category(cat["source_name"], masters)
		cat["match"], cat["candidates"] = best, candidates
		if not best:
			options = ", ".join(f"{c['category_name']} ({c['score']:.0%})" for c in candidates) or "none"
			errors.append(f"Row {cat['row']}: category '{cat['source_name'][:60]}' has no master ≥ 80% "
						  f"— closest: {options}.")
			continue
		if best["name"] in resolved:
			errors.append(f"Rows {resolved[best['name']]['row']} and {cat['row']}: '{resolved[best['name']]['source_name'][:40]}' "
						  f"and '{cat['source_name'][:40]}' both match category '{best['category_name']}' "
						  "(a category may appear once per budget).")
		else:
			resolved[best["name"]] = cat
